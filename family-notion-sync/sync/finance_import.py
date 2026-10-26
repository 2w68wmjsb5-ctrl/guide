"""Parse transaction exports (CSV/XLSX) from finance apps and bank portals.

Taken from the family-dashboard app. Exports differ per app, so columns are
matched by common German/English header names instead of a fixed format.
"""

import csv
import hashlib
import io
import logging
import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

log = logging.getLogger(__name__)

HEADER_ALIASES = {
    "date": ["buchungstag", "buchungsdatum", "datum", "date", "wertstellung", "valutadatum",
             "valuta", "booking date", "transaction date"],
    "amount": ["betrag", "betrag (eur)", "betrag in eur", "betrag (€)", "umsatz", "amount",
               "betrag eur", "umsatz in eur"],
    "debit": ["soll", "ausgang", "debit"],
    "credit": ["haben", "eingang", "credit"],
    "counterparty": ["empfänger", "empfaenger", "auftraggeber/empfänger", "name zahlungsbeteiligter",
                     "beguenstigter/zahlungspflichtiger", "begünstigter/zahlungspflichtiger",
                     "beguenstigter/auftraggeber", "begünstigter/auftraggeber", "zahlungsempfänger",
                     "zahlungspflichtige*r", "zahlungsempfänger*in", "gegenseite", "payee", "name",
                     "partner", "auftraggeber / begünstigter"],
    "purpose": ["verwendungszweck", "buchungstext", "beschreibung", "description", "vorgang",
                "notiz", "reference"],
    "category": ["kategorie", "category", "hauptkategorie", "analyse-hauptkategorie"],
    "account": ["konto", "kontoname", "account", "referenzkonto", "iban auftragskonto"],
}


AMEX = "Amex"
# Columns only American Express exports have (German and English online banking)
AMEX_HEADERS = {"erscheint auf ihrer abrechnung als", "appears on your statement as",
                "erweiterte details", "extended details"}
# Your own payments towards the card bill, not spending
AMEX_PAYMENT = re.compile(r"zahlung erhalten|besten dank|payment received|thank you|lastschrifteinzug", re.I)


@dataclass
class Transaction:
    day: date
    amount: float
    counterparty: str
    purpose: str
    category: str = ""
    account: str = ""

    @property
    def uid(self) -> str:
        raw = f"{self.day}|{self.amount:.2f}|{self.counterparty}|{self.purpose}|{self.account}"
        return hashlib.sha1(raw.encode()).hexdigest()


def _norm(header: str) -> str:
    return re.sub(r"\s+", " ", str(header or "").replace("﻿", "")).strip().strip('"').lower()


def map_headers(headers: list[str]) -> dict[str, int]:
    normed = [_norm(h) for h in headers]
    mapping: dict[str, int] = {}
    for field, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            if alias in normed:
                mapping[field] = normed.index(alias)
                break
    return mapping


def parse_amount(value) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace("€", "").replace("EUR", "").replace("\xa0", "").replace(" ", "").strip()
    if not text:
        return None
    negative = text.startswith("-") or text.endswith("-")
    text = text.strip("+-")
    if "," in text and "." in text:
        # Whichever separator comes last is the decimal separator
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        number = float(text)
    except ValueError:
        return None
    return -number if negative else number


def parse_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%Y-%m-%dT%H:%M:%S"):
        try:
            return datetime.strptime(text[:19] if "T" in fmt else text, fmt).date()
        except ValueError:
            continue
    return None


def _rows_from_csv(raw: bytes) -> list[list[str]]:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    sample = "\n".join(text.splitlines()[:30])
    delimiter = max([";", ",", "\t"], key=sample.count)
    return list(csv.reader(io.StringIO(text), delimiter=delimiter))


def _rows_from_xlsx(path: Path) -> list[list]:
    from openpyxl import load_workbook

    sheet = load_workbook(path, read_only=True, data_only=True).active
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def parse_rows(rows: list[list], filename: str = "") -> list[Transaction]:
    """Find the header line (bank exports often have a preamble) and parse rows below it."""
    for idx, row in enumerate(rows[:40]):
        headers = [str(c) if c is not None else "" for c in row]
        mapping = map_headers(headers)
        has_amount = "amount" in mapping or ("debit" in mapping and "credit" in mapping)
        if "date" in mapping and has_amount:
            break
    else:
        raise ValueError("Keine Spalten für Datum und Betrag gefunden")
    # Amex lists purchases as positive amounts and also contains the card payments
    is_amex = bool(AMEX_HEADERS & {_norm(h) for h in headers}) or "amex" in filename.lower()

    def cell(row, field):
        i = mapping.get(field)
        return row[i] if i is not None and i < len(row) and row[i] is not None else ""

    result = []
    for row in rows[idx + 1:]:
        day = parse_date(cell(row, "date"))
        if "amount" in mapping:
            amount = parse_amount(cell(row, "amount"))
        else:
            debit = parse_amount(cell(row, "debit")) or 0.0
            credit = parse_amount(cell(row, "credit")) or 0.0
            amount = credit - abs(debit) if (debit or credit) else None
        if day is None or amount is None:
            continue
        counterparty = str(cell(row, "counterparty")).strip()
        purpose = str(cell(row, "purpose")).strip()
        account = str(cell(row, "account")).strip()
        if is_amex:
            if amount < 0 and AMEX_PAYMENT.search(f"{counterparty} {purpose}"):
                continue
            amount, account = -amount, AMEX
        result.append(Transaction(
            day=day,
            amount=round(amount, 2),
            counterparty=counterparty,
            purpose=purpose,
            category=str(cell(row, "category")).strip(),
            account=account,
        ))
    return result


def parse_file(path: Path) -> list[Transaction]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        return parse_rows(_rows_from_xlsx(path), path.name)
    return parse_rows(_rows_from_csv(path.read_bytes()), path.name)
